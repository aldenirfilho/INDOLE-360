from pathlib import Path
import json,re
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image,ImageDraw,ImageFont
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'site/documentos';OUT.mkdir(exist_ok=True)
D=json.loads((ROOT/'site/data/catalogo.json').read_text());C=D['companies']
def new(title,subtitle=None):
 d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.7);sec.bottom_margin=Cm(1.6);sec.left_margin=sec.right_margin=Cm(1.8)
 for border in d.styles.element.xpath('.//w:pBdr'):border.getparent().remove(border)
 for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
  st=d.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
 d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.space_after=Pt(6);d.styles['Normal'].paragraph_format.line_spacing=1.08
 d.styles['Title'].font.size=Pt(25);d.styles['Heading 1'].font.size=Pt(17);d.styles['Heading 2'].font.size=Pt(13);d.styles['Heading 3'].font.size=Pt(11)
 d.styles['Heading 2'].paragraph_format.space_before=Pt(12)
 d.core_properties.author='Aldenir Rocha de Oliveira Filho';d.core_properties.title=title
 d.add_paragraph(title,'Title')
 if subtitle:d.add_paragraph(subtitle,'Subtitle')
 foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT;r=foot.add_run('ÍNDOLE 360  |  v2.0  |  ');r.font.size=Pt(8)
 f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
 return d
def p(d,t,style=None):return d.add_paragraph(t,style)
def h(d,t,level=1):return d.add_heading(re.sub('[#°:—–×]',' ',t).strip(),level=level)
def hyperlink(par,text,url):
 if not url:return
 part=par.part;rid=part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True);el=OxmlElement('w:hyperlink');el.set(qn('r:id'),rid);r=OxmlElement('w:r');rp=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'225A46');rp.append(co);r.append(rp);t=OxmlElement('w:t');t.text=text;r.append(t);el.append(r);par._p.append(el)
def table(d,headers,rows,widths=None):
 t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths:
  for col,w in zip(t.columns,widths):col.width=Cm(w)
 for cell,text in zip(t.rows[0].cells,headers):cell.text=text
 hdr=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(hdr)
 for row in rows:
  for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
 for i,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();cs=OxmlElement('w:cantSplit');trpr.append(cs)
  for j,cell in enumerate(row.cells):
   if widths:cell.width=Cm(widths[j])
   cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   pr.append(b);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'173C35' if i==0 else ('F0F4EC' if i%2==0 else 'FFFFFF'));pr.append(sh)
   margins=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');margins.append(e)
   pr.append(margins)
   for pa in cell.paragraphs:
    pa.paragraph_format.space_after=Pt(2);pa.paragraph_format.space_before=Pt(2)
    for r in pa.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000');r.bold=i==0
 d.add_paragraph().paragraph_format.space_after=Pt(2)
 return t
def md(d,text,skiptitle=False):
 lines=text.splitlines();i=0
 while i<len(lines):
  s=lines[i].strip();i+=1
  if not s:continue
  if s.startswith('|'):
   rr=[s]
   while i<len(lines) and lines[i].strip().startswith('|'):rr.append(lines[i]);i+=1
   rows=[[z.strip().replace('**','') for z in x.strip().strip('|').split('|')] for x in rr if not re.match(r'^\|[\s:|\-]+\|$',x)]
   if len(rows)>1:table(d,rows[0],rows[1:])
   continue
  if s.startswith('#'):
   level=len(s)-len(s.lstrip('#'))
   if skiptitle and level==1:continue
   h(d,s.lstrip('# '),min(max(level-1,1),3));continue
  s=s.replace('**','').replace('`','').removeprefix('> ')
  s=re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1',s)
  p(d,s)
def fmt(m):
 if m.get('value') is None:return 'N/D'
 v=f"{m['value']:,.2f}".rstrip('0').rstrip('.').replace(',','X').replace('.',',').replace('X','.')
 return (m.get('currency') or '')+' '+v+(' mi' if m.get('unit')=='million' else '')
def first_sentence(s,maxlen=330):
 if len(s)<=maxlen:return s
 pieces=re.split(r'(?<=[.!?])\s+',s);out=''
 for a in pieces:
  if len(out)+len(a)>maxlen:break
  out+=(' ' if out else '')+a
 return out or s
def source_line(d,c):
 pa=p(d,'Fontes consultadas em 18/09/2026: ')
 for i,s in enumerate(c['sources'],1):
  hyperlink(pa,f'[{i}] {s.get("title") or s.get("id") or "Documento"}',s['url']);pa.add_run('  ')
 for r in pa.runs:r.font.size=Pt(8)
# Exact explanatory visual; no implied observed allocation.
im=Image.new('RGB',(1800,380),'#f6f4eb');dr=ImageDraw.Draw(im);font=ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf',40);small=ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf',28);left=30
for v,label,col in zip([30,30,30,10],['Produto','Pessoas','Impacto','Reserva'],['#2e6150','#739175','#b48846','#dedbce']):
 width=17.4*v;dr.rounded_rectangle((left,30,left+width-5,265),radius=12,fill=col);dr.text((left+width/2,102),str(v)+'%',font=font,fill='white' if v!=10 else '#173c35',anchor='mm');dr.text((left+width/2,165),label,font=small,fill='white' if v!=10 else '#173c35',anchor='mm');left+=width
dr.text((900,330),'PROPOSTA AUTORAL · Não representa gastos observados das empresas',font=small,fill='#173c35',anchor='mm');im.save(ROOT/'site/assets/modelo-30303010.png')
doc=new('ÍNDOLE 360 metodologia e projetos','Modelo autoral 30 30 30 10  |  Aldenir Rocha  |  18 setembro 2026')
p(doc,'Este documento explica o propósito do catálogo, sua régua de avaliação, os limites contábeis e os projetos de humanização e simplificação tecnológica. A conclusão operacional é iniciar com evidências rastreáveis e pilotos avaliáveis, sem atribuir notas onde faltam dados.')
doc.add_picture(str(ROOT/'site/assets/modelo-30303010.png'),width=Cm(17.2))
h(doc,'Roteiro de leitura')
p(doc,'1  Propósito e modelo  •  2  Contabilidade  •  3  Indicadores  •  4  Agentes  •  5  Humanização  •  6  Cenários  •  7  Integridade  •  8  Aperfeiçoamento')
md(doc,(ROOT/'docs/METODOLOGIA.md').read_text(),True)
h(doc,'Projeto separado Humanização da Tecnologia').paragraph_format.page_break_before=True;md(doc,(ROOT/'docs/HUMANIZACAO_TECNOLOGIA.md').read_text(),True)
h(doc,'Estratégia de adoção e validação').paragraph_format.page_break_before=True;md(doc,(ROOT/'docs/NUCLEO_ESTRATEGICO_ATOMIC.md').read_text(),True)
doc.save(OUT/'01_Metodologia_e_Projetos.docx')
# Build a fresh catalog to ensure no placeholder pages survive.
doc=new('ÍNDOLE 360 catálogo das empresas','Dez grandes empresas  |  Base 17 setembro 2026  |  v2.0')
p(doc,'Amostra selecionada por valor de mercado, não por conduta. Dez fichas documentais com informações financeiras, impactos e sugestões. Notas gerais e aderência 30/30/30/10 ainda não concluídas; faltam conciliação de recursos e revisão das onze dimensões.')
table(doc,['Empresa','Setor','Exercício'],[[c['short_name'],c['sector'],c['fiscal_label']] for c in C],[4,9.4,4])
p(doc,'mi = milhões; USD = dólar dos Estados Unidos; TWD = dólar de Taiwan; P&D = pesquisa e desenvolvimento. Os períodos fiscais diferem. Ausência de dado não significa ausência de ação. Fontes corporativas não substituem verificação independente de impacto.')
hyperlink(p(doc,''),'Seleção CompaniesMarketCap consultada em 18/09/2026 com fotografia de 17/09/2026',D['selection']['source'])
hyperlink(p(doc,''),'Site com fichas completas e fontes por campo','https://aldenirfilho.github.io/INDOLE-360/')
for c in C:
 doc.add_page_break();h(doc,c['short_name']);p(doc,c['sector']+' | '+c['country']+' | '+' · '.join(t['symbol'] for t in c['tickers']))
 p(doc,c['business_model']);p(doc,c['assessment_summary'])
 table(doc,['Dado','Valor','Período'],[[label,fmt(c[k]),c['fiscal_label'] if k!='employees' else c[k]['period']] for k,label in [('revenue','Receita'),('net_income','Lucro líquido'),('research','P&D'),('capex','CAPEX'),('dividends','Dividendos'),('buybacks','Recompras'),('employees','Empregados')]], [4,5.4,8])
 p(doc,'Rubricas não somáveis. '+(c['net_income'].get('note') or '')+' '+(c['employees'].get('note') or c['employees'].get('definition') or ''))
 h(doc,'Impacto trabalho e aplicações',2)
 for label,items in [('Social e ambiente',c['social_environment']),('Trabalho',c['worker']),('Aplicações',c['investment_evidence'])]:
  if items:
   e=items[0];text=label+': '+first_sentence(e['claim'],260)
   if e.get('limitation'):text+=' '+first_sentence(e['limitation'],150)
   pa=p(doc,text);hyperlink(pa,' [fonte]',e.get('source'))
 h(doc,'Prioridades de melhoria',2)
 for x in c['improvements']:p(doc,x,'List Bullet')
 # Legal case and complete primary sources have their own page: preserves exact status.
 doc.add_page_break();h(doc,c['short_name']+' evidências e limites')
 h(doc,'Casos selecionados e situação conhecida',2)
 if not c['cases']:p(doc,'Nenhum caso com desfecho suficiente foi incluído nesta coleta. Isso não equivale à ausência de processos, danos ou irregularidades.')
 for e in c['cases']:
  pa=p(doc,(e['claim'] or '')+' | '+(e.get('verification') or '')+' | '+str(e.get('period') or 'Data não informada'));p(doc,e.get('text') or 'Estado não confirmado')
  if e.get('limitation'):p(doc,e['limitation'])
  hyperlink(p(doc,''),'Documento primário',e.get('source'))
 h(doc,'Lacunas que impedem a nota geral',2)
 for x in c['gaps'][:5]:p(doc,x,'List Bullet')
 h(doc,'Fontes e documentação',2);source_line(doc,c)
 p(doc,'Os dados financeiros, as práticas anunciadas e os impactos não possuem necessariamente o mesmo nível de verificação. Consulta documental não equivale a auditoria independente da empresa. A íntegra por campo e as ressalvas adicionais estão no catálogo digital.')
 hyperlink(p(doc,''),'Abrir ficha completa',f'https://aldenirfilho.github.io/INDOLE-360/#empresa={c["id"]}')
doc.save(OUT/'02_Catalogo_das_Empresas.docx')
doc=new('Cartas de melhoria empresarial','ÍNDOLE 360  |  Aldenir Rocha  |  18 setembro 2026')
p(doc,'Dez cartas construtivas para liderança e equipes responsáveis. A versão em português permite revisão e edição; arquivos em inglês acompanham o pacote para correspondência internacional. Canais institucionais não garantem encaminhamento ao CEO. O registro de envio e entrega é separado destas cartas.')
for c in C:
 doc.add_page_break();h(doc,'Carta para '+c['short_name']);txt=(ROOT/'site/cartas'/f'{c["id"]}.txt').read_text()
 for para in txt.split('\n\n')[1:]:p(doc,para)
 pc=c['public_contact'];p(doc,'Canal público: '+str(pc.get('email') or 'E-mail não confirmado; consultar canal oficial.'));hyperlink(p(doc,''),'Verificação do contato institucional',pc['source'])
doc.save(OUT/'03_Cartas_de_Melhoria.docx')
print('\n'.join(str(p) for p in OUT.glob('*.docx')))
