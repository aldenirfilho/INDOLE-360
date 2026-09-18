from pathlib import Path
import json,re
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image,ImageDraw,ImageFont
ROOT=Path(__file__).resolve().parents[1]
def new(title,subtitle=None):
 d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.7);sec.bottom_margin=Cm(2.0);sec.footer_distance=Cm(0.7);sec.left_margin=sec.right_margin=Cm(1.8)
 for border in d.styles.element.xpath('.//w:pBdr'):border.getparent().remove(border)
 for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
  st=d.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
 d.styles['Normal'].font.size=Pt(10.5);d.styles['Normal'].paragraph_format.space_after=Pt(6);d.styles['Normal'].paragraph_format.line_spacing=1.08
 d.styles['Title'].font.size=Pt(25);d.styles['Heading 1'].font.size=Pt(17);d.styles['Heading 2'].font.size=Pt(13);d.styles['Heading 3'].font.size=Pt(11)
 d.styles['Heading 2'].paragraph_format.space_before=Pt(12)
 d.core_properties.author='Aldenir Rocha de Oliveira Filho';d.core_properties.title=title
 d.add_paragraph(title,'Title')
 if subtitle:d.add_paragraph(subtitle,'Subtitle')
 foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT;r=foot.add_run('ÍNDOLE 360  |  v2.1  |  ');r.font.size=Pt(8)
 f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
 return d

def p(d,t,style=None):return d.add_paragraph(t,style)

def h(d,t,level=1):return d.add_heading(re.sub('[#°:—–×]',' ',t).strip(),level=level)

def hyperlink(par,text,url):
 if not url:return
 part=par.part;rid=part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True);el=OxmlElement('w:hyperlink');el.set(qn('r:id'),rid);r=OxmlElement('w:r');rp=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'225A46');rp.append(co);r.append(rp);t=OxmlElement('w:t');t.text=text;r.append(t);el.append(r);par._p.append(el)

def table(d,headers,rows,widths=None):
 t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths:
  for col,w in zip(t.columns,widths):col.width=Cm(w)
 for cell,text in zip(t.rows[0].cells,headers):cell.text=text
 hdr=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(hdr)
 for row in rows:
  for cell,text in zip(t.add_row().cells,row):cell.text=str(text)
 for i,row in enumerate(t.rows):
  trpr=row._tr.get_or_add_trPr();cs=OxmlElement('w:cantSplit');trpr.append(cs)
  for j,cell in enumerate(row.cells):
   if widths:cell.width=Cm(widths[j])
   cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   pr.append(b);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'173C35' if i==0 else ('F0F4EC' if i%2==0 else 'FFFFFF'));pr.append(sh)
   margins=OxmlElement('w:tcMar')
   for edge in ['top','bottom','left','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');margins.append(e)
   pr.append(margins)
   for pa in cell.paragraphs:
    pa.paragraph_format.space_after=Pt(2);pa.paragraph_format.space_before=Pt(2)
    for r in pa.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000');r.bold=i==0
 d.add_paragraph().paragraph_format.space_after=Pt(2)
 return t

def inline(par,text):
 pattern=r'(\[([^\]]+)\]\(([^)]+)\)|\*\*([^*]+)\*\*)'
 pos=0
 for m in re.finditer(pattern,text):
  par.add_run(text[pos:m.start()])
  if m.group(2):
   url=m.group(3)
   if not url.startswith('https://'):
    path=(ROOT/'docs'/url).resolve().relative_to(ROOT).as_posix()
    url='https://github.com/aldenirfilho/INDOLE-360/blob/main/'+path
   hyperlink(par,m.group(2),url)
  else:par.add_run(m.group(4)).bold=True
  pos=m.end()
 par.add_run(text[pos:])

def markdown(d,text):
 lines=text.splitlines();i=0
 while i<len(lines):
  s=lines[i].strip();i+=1
  if not s or s.startswith('# '):continue
  if s.startswith('|'):
   rr=[s]
   while i<len(lines) and lines[i].strip().startswith('|'):rr.append(lines[i]);i+=1
   rows=[[z.strip().replace('**','') for z in x.strip().strip('|').split('|')] for x in rr if not re.match(r'^\|[\s:|\-]+\|$',x)]
   table(d,rows[0],rows[1:]);continue
  if s.startswith('## '):
   h(d,s[3:],1);continue
  style='List Bullet' if s.startswith('- ') else None
  if style:s=s[2:]
  inline(d.add_paragraph(style=style),s)

doc=new('Equalização e contribuição proporcional','ÍNDOLE 360°  |  Complemento v2.1  |  18 setembro 2026')
p(doc,'Roteiro: 1. Propósito e cálculo  •  2. Portfólio e veracidade  •  3. Índole dos Impostos  •  4. Revisão cruzada e próximos passos')
doc.add_picture(str(ROOT/'site/assets/modelo-30303010.png'),width=Cm(17.2))
markdown(doc,(ROOT/'docs/EQUALIZACAO_E_CONTRIBUICAO.md').read_text())
h(doc,'Índole dos Impostos concepção inicial');markdown(doc,(ROOT/'docs/INDOLE_DOS_IMPOSTOS.md').read_text())
doc.add_page_break();h(doc,'Revisão cruzada GPT e Codex');markdown(doc,(ROOT/'docs/REVISAO_CRUZADA_GPT_CODEX.md').read_text())
# Keep title rules absent across Word and Google Docs renderers.
for border in doc.element.xpath('.//w:pBdr'):border.getparent().remove(border)
out=ROOT/'site/documentos/04_Equalizacao_e_Contribuicao.docx';doc.save(out);print(out)
